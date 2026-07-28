"""قراءة ملف النتيجة بأي صيغة وتحويله لسجلات (dict لكل طالب).

الصيغ المدعومة: CSV / TSV / TXT / XLSX / XLSM / XLS / JSON / JSONL / SQLite.
"""

from __future__ import annotations

import codecs
import csv
import itertools
import json
import sqlite3
import sys
from pathlib import Path
from typing import Iterable, Iterator

from .schema import detect_fields

HEADER_SCAN_ROWS = 15
DELIMITED_EXT = {".csv", ".tsv", ".txt", ".dat"}
EXCEL_EXT = {".xlsx", ".xlsm"}
LEGACY_EXCEL_EXT = {".xls"}
JSON_EXT = {".json"}
JSONL_EXT = {".jsonl", ".ndjson"}
SQLITE_EXT = {".db", ".sqlite", ".sqlite3"}

SUPPORTED_EXT = (
    DELIMITED_EXT | EXCEL_EXT | LEGACY_EXCEL_EXT | JSON_EXT | JSONL_EXT | SQLITE_EXT
)

csv.field_size_limit(min(sys.maxsize, 2**31 - 1))


class ReaderError(RuntimeError):
    """خطأ في قراءة الملف نعرضه للمستخدم مباشرة."""


def read_records(
    path: str | Path,
    sheet: str | None = None,
    table: str | None = None,
    header_row: int | None = None,
) -> Iterator[dict]:
    """يرجّع سجلات الطلبة كـ dict {عنوان العمود: القيمة}."""
    path = Path(path)
    if not path.exists():
        raise ReaderError(f"الملف مش موجود: {path}")

    ext = path.suffix.lower()
    if ext in JSON_EXT:
        yield from _read_json(path)
    elif ext in JSONL_EXT:
        yield from _read_jsonl(path)
    elif ext in SQLITE_EXT:
        yield from _read_sqlite(path, table)
    elif ext in EXCEL_EXT:
        yield from _rows_to_records(_read_xlsx(path, sheet), header_row)
    elif ext in LEGACY_EXCEL_EXT:
        yield from _rows_to_records(_read_xls(path, sheet), header_row)
    elif ext in DELIMITED_EXT:
        yield from _rows_to_records(_read_delimited(path), header_row)
    else:
        raise ReaderError(
            f"صيغة غير مدعومة «{ext}». الصيغ المدعومة: "
            + "، ".join(sorted(SUPPORTED_EXT))
        )


# ---------------------------------------------------------------- الملفات النصية


def _detect_encoding(path: Path) -> str:
    with open(path, "rb") as handle:
        sample = handle.read(1 << 16)
    if sample.startswith(codecs.BOM_UTF8):
        return "utf-8-sig"
    if sample.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return "utf-16"
    try:
        sample.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError as exc:
        # لو الخطأ في آخر البايتات يبقى حرف مقطوع مش ترميز غلط
        if exc.start >= len(sample) - 4:
            return "utf-8"
    return "cp1256"


def _detect_delimiter(path: Path, encoding: str) -> str:
    with open(path, encoding=encoding, errors="replace", newline="") as handle:
        sample = handle.read(8192)
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return "\t" if path.suffix.lower() in {".tsv", ".txt", ".dat"} else ","


def _read_delimited(path: Path) -> Iterator[list]:
    encoding = _detect_encoding(path)
    delimiter = _detect_delimiter(path, encoding)
    with open(path, encoding=encoding, errors="replace", newline="") as handle:
        for row in csv.reader(handle, delimiter=delimiter):
            yield row


# ---------------------------------------------------------------------- إكسيل


def _read_xlsx(path: Path, sheet: str | None) -> Iterator[list]:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - يعتمد على البيئة
        raise ReaderError(
            "محتاج تثبّت openpyxl عشان تقرأ ملفات xlsx:  pip install openpyxl"
        ) from exc

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet:
            if sheet not in workbook.sheetnames:
                raise ReaderError(
                    f"الشيت «{sheet}» مش موجود. الشيتات المتاحة: "
                    + "، ".join(workbook.sheetnames)
                )
            worksheet = workbook[sheet]
        else:
            worksheet = workbook.worksheets[0]
        for row in worksheet.iter_rows(values_only=True):
            yield list(row)
    finally:
        workbook.close()


def _read_xls(path: Path, sheet: str | None) -> Iterator[list]:
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover - يعتمد على البيئة
        raise ReaderError(
            "محتاج تثبّت xlrd عشان تقرأ ملفات xls القديمة:  pip install xlrd"
        ) from exc

    book = xlrd.open_workbook(path)
    worksheet = book.sheet_by_name(sheet) if sheet else book.sheet_by_index(0)
    for index in range(worksheet.nrows):
        yield worksheet.row_values(index)


# ----------------------------------------------------------------------- JSON


def _iter_json_records(payload) -> Iterator[dict]:
    if isinstance(payload, dict):
        # ممكن يكون {"data": [...]} أو {"رقم الجلوس": {...}}
        for key in ("data", "results", "students", "rows", "items"):
            if isinstance(payload.get(key), list):
                payload = payload[key]
                break
        else:
            for key, value in payload.items():
                if isinstance(value, dict):
                    record = {"رقم الجلوس": key}
                    record.update(value)
                    yield record
            return
    if not isinstance(payload, list):
        raise ReaderError("شكل ملف JSON مش مفهوم — المفروض يكون قائمة سجلات.")
    for item in payload:
        if isinstance(item, dict):
            yield item


def _read_json(path: Path) -> Iterator[dict]:
    encoding = _detect_encoding(path)
    with open(path, encoding=encoding, errors="replace") as handle:
        try:
            payload = json.load(handle)
        except json.JSONDecodeError as exc:
            raise ReaderError(f"ملف JSON فيه خطأ: {exc}") from exc
    yield from _iter_json_records(payload)


def _read_jsonl(path: Path) -> Iterator[dict]:
    encoding = _detect_encoding(path)
    with open(path, encoding=encoding, errors="replace") as handle:
        for line_no, line in enumerate(handle, start=1):
            line = line.strip()
            if not line:
                continue
            try:
                item = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ReaderError(f"سطر {line_no} في ملف JSONL فيه خطأ: {exc}") from exc
            if isinstance(item, dict):
                yield item


# --------------------------------------------------------------------- SQLite


def _read_sqlite(path: Path, table: str | None) -> Iterator[dict]:
    connection = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    try:
        if not table:
            tables = [
                row[0]
                for row in connection.execute(
                    "SELECT name FROM sqlite_master WHERE type='table' "
                    "AND name NOT LIKE 'sqlite_%'"
                )
            ]
            if not tables:
                raise ReaderError("مفيش جداول في قاعدة البيانات دي.")
            if len(tables) > 1:
                raise ReaderError(
                    "فيه أكتر من جدول، حدّد واحد بـ --table. الجداول: "
                    + "، ".join(tables)
                )
            table = tables[0]
        quoted = '"' + table.replace('"', '""') + '"'
        for row in connection.execute(f"SELECT * FROM {quoted}"):
            yield dict(row)
    finally:
        connection.close()


# ------------------------------------------------------ تحويل صفوف خام لسجلات


def _clean_header(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())


def _pick_header_index(rows: list[list]) -> int:
    """بيدوّر على صف العناوين — بعض الملفات بتبدأ بسطور عنوان أو ترويسة."""
    best_index, best_score = -1, 0
    for index, row in enumerate(rows):
        score = len(detect_fields([_clean_header(cell) for cell in row]))
        if score > best_score:
            best_index, best_score = index, score
    if best_score >= 2:
        return best_index
    for index, row in enumerate(rows):
        if sum(1 for cell in row if _clean_header(cell)) >= 2:
            return index
    return 0


def _rows_to_records(rows: Iterable[list], header_row: int | None) -> Iterator[dict]:
    iterator = iter(rows)
    buffered: list[list] = []
    for row in iterator:
        buffered.append(list(row))
        if len(buffered) >= HEADER_SCAN_ROWS:
            break
    if not buffered:
        return

    if header_row is not None:
        index = header_row - 1
        if index < 0 or index >= len(buffered):
            raise ReaderError(
                f"--header-row لازم يكون بين 1 و{len(buffered)} حسب أول صفوف الملف."
            )
    else:
        index = _pick_header_index(buffered)

    headers = [_clean_header(cell) for cell in buffered[index]]
    if not any(headers):
        raise ReaderError("مقدرتش أحدد صف العناوين — جرّب تحدده بـ --header-row.")

    for row in itertools.chain(buffered[index + 1 :], iterator):
        if not any(str(cell).strip() for cell in row if cell is not None):
            continue
        record = {}
        for position, header in enumerate(headers):
            if not header:
                continue
            record[header] = row[position] if position < len(row) else None
        yield record
